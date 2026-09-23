"""Scrub pipeline tests: full-row result files for CSV and XLSX uploads (DNC lookup mocked)."""
import csv
import io
import tempfile
from unittest import mock

import openpyxl
from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.test import TestCase, override_settings

from .dnc import BatchResult
from .models import ScrubJob
from .tasks import run_scrub_job

User = get_user_model()
_MEDIA = tempfile.mkdtemp()


def _fake_run_checks(numbers, scrub_types, stop_event=None):
    """Numbers starting with 903 are DNC, everything else is clean."""
    dnc = [n for n in numbers if n.startswith('903')]
    clean = [n for n in numbers if not n.startswith('903')]
    return BatchResult(clean=clean, dnc_numbers=dnc, unchecked=[])


@override_settings(
    MEDIA_ROOT=_MEDIA,
    DEFAULT_FILE_STORAGE='django.core.files.storage.FileSystemStorage',
    CACHES={'default': {'BACKEND': 'django.core.cache.backends.locmem.LocMemCache'}},
    EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend',
)
@mock.patch('scrubber.tasks.run_checks', side_effect=_fake_run_checks)
class ResultFileTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(email='u@example.com', password='x', credits=1000)

    def _job(self, name, payload):
        job = ScrubJob(user=self.user, filename=name, scrub_types=['federal_dnc'],
                       status=ScrubJob.Status.QUEUED)
        job.file.save(name, ContentFile(payload), save=True)
        return job

    def test_csv_keeps_all_columns(self, _rc):
        payload = (
            'CELLPHONE,FIRST,LAST,CITY\n'
            '9032757138,Georgia,Kirvy,Wills Point\n'
            '4405829719,Frank,Toth,North Royalton\n'
            '(440) 582-9719,Dup,Row,Nowhere\n'
            'bad,No,Phone,Here\n'
        ).encode()
        job = self._job('leads.csv', payload)
        run_scrub_job(job.pk)
        job.refresh_from_db()

        self.assertEqual(job.status, ScrubJob.Status.COMPLETED)
        self.assertEqual((job.total, job.clean, job.dnc), (2, 1, 1))
        self.assertTrue(job.result_file.name.endswith('_clean.csv'))

        clean = list(csv.reader(io.StringIO(job.result_file.read().decode('utf-8-sig'))))
        self.assertEqual(clean, [
            ['CELLPHONE', 'FIRST', 'LAST', 'CITY', 'dnc_status'],
            ['4405829719', 'Frank', 'Toth', 'North Royalton', 'Clean'],
        ])
        dnc = list(csv.reader(io.StringIO(job.result_file_dnc.read().decode('utf-8-sig'))))
        self.assertEqual(dnc[1], ['9032757138', 'Georgia', 'Kirvy', 'Wills Point', 'Do Not Call'])
        self.user.refresh_from_db()
        self.assertEqual(self.user.credits, 998)

    def test_xlsx_in_xlsx_out(self, _rc):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['Name', 'Phone', 'Zip'])
        ws.append(['Ann', 9032757138, 75169])
        ws.append(['Bob', 4405829719.0, 44133])
        buf = io.BytesIO()
        wb.save(buf)
        job = self._job('leads.xlsx', buf.getvalue())
        run_scrub_job(job.pk)
        job.refresh_from_db()

        self.assertEqual(job.status, ScrubJob.Status.COMPLETED)
        self.assertEqual((job.total, job.clean, job.dnc), (2, 1, 1))
        self.assertTrue(job.result_file.name.endswith('_clean.xlsx'))
        self.assertTrue(job.result_file_dnc.name.endswith('_dnc.xlsx'))

        out = openpyxl.load_workbook(io.BytesIO(job.result_file.read())).active
        rows = [list(r) for r in out.iter_rows(values_only=True)]
        self.assertEqual(rows, [['Name', 'Phone', 'Zip', 'dnc_status'], ['Bob', 4405829719, 44133, 'Clean']])
        out = openpyxl.load_workbook(io.BytesIO(job.result_file_dnc.read())).active
        rows = [list(r) for r in out.iter_rows(values_only=True)]
        self.assertEqual(rows, [['Name', 'Phone', 'Zip', 'dnc_status'], ['Ann', 9032757138, 75169, 'Do Not Call']])

    def test_plain_number_list_still_works(self, _rc):
        job = self._job('numbers.txt', b'(903) 275-7138\n440-582-9719\n')
        run_scrub_job(job.pk)
        job.refresh_from_db()
        clean = job.result_file.read().decode('utf-8-sig').splitlines()
        self.assertEqual(clean, ['phone_number,dnc_status', '440-582-9719,Clean'])

    def test_resume_after_pause_keeps_columns(self, _rc):
        payload = b'phone,name\n9032757138,A\n4405829719,B\n'
        job = self._job('leads.csv', payload)
        # Simulate a paused job: B still remaining, A already flagged DNC.
        job.total = 2
        job.status = ScrubJob.Status.PAUSED
        job.partial_data_file.save(
            'p.json',
            ContentFile(b'{"clean": [], "dnc": ["9032757138"], "remaining": ["4405829719"]}'),
            save=True,
        )
        run_scrub_job(job.pk)
        job.refresh_from_db()
        self.assertEqual(job.status, ScrubJob.Status.COMPLETED)
        clean = list(csv.reader(io.StringIO(job.result_file.read().decode('utf-8-sig'))))
        self.assertEqual(clean, [['phone', 'name', 'dnc_status'], ['4405829719', 'B', 'Clean']])
        dnc = list(csv.reader(io.StringIO(job.result_file_dnc.read().decode('utf-8-sig'))))
        self.assertEqual(dnc, [['phone', 'name', 'dnc_status'], ['9032757138', 'A', 'Do Not Call']])


@override_settings(
    MEDIA_ROOT=_MEDIA,
    DEFAULT_FILE_STORAGE='django.core.files.storage.FileSystemStorage',
)
class UploadViewTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(email='v@example.com', password='x', credits=10)
        self.client.force_login(self.user)

    def _post(self, name, payload):
        from django.core.files.uploadedfile import SimpleUploadedFile
        with mock.patch('scrubber.tasks.process_scrub_job.delay') as delay:
            resp = self.client.post(
                '/scrubber/',
                {'file': SimpleUploadedFile(name, payload), 'scrub_types': ['federal_dnc']},
                HTTP_X_REQUESTED_WITH='XMLHttpRequest',
            )
        return resp, delay

    def test_xlsx_accepted(self):
        resp, delay = self._post('leads.xlsx', b'PK\x03\x04fake')
        self.assertEqual(resp.status_code, 200, resp.content)
        self.assertTrue(resp.json()['ok'])
        self.assertTrue(delay.called)

    def test_xls_rejected(self):
        resp, _ = self._post('leads.xls', b'old')
        self.assertEqual(resp.status_code, 400)
        self.assertIn('.xlsx', resp.json()['error'])
