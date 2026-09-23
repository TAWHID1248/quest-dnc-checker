"""
Phone number normalization and upload-file parsing.

All numbers are normalised to bare 10-digit US format (NPANXXXXXX) before
any DNC checking occurs.

Parsing keeps every column of every row so the result files can hand the
user back their full records (name, address, e-mail, …) split into the
clean and DNC sets, not just the phone numbers.

Supported uploads: .csv / .txt (any common delimiter) and .xlsx (first sheet).
This module has no Django dependency; openpyxl is imported lazily for .xlsx.
"""

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Iterator

# Pre-compiled once at import time
_DIGIT_RE   = re.compile(r'\D')
_VALID_RE   = re.compile(r'^[2-9]\d{2}[2-9]\d{6}$')  # NANP: NPA & NXX can't start with 0 or 1
_ALPHA_RE   = re.compile(r'[A-Za-z]')

# Header names that identify the phone column (checked case-insensitively).
_PHONE_HEADER_RE = re.compile(
    r'phone|cell|mobile|tel|number|contact', re.IGNORECASE,
)

EXCEL_EXTENSIONS = ('.xlsx', '.xlsm')


def normalize(raw: str) -> str | None:
    """
    Strip a raw phone string down to 10 clean US digits.

    Returns the 10-digit string on success, or None if the number is
    structurally invalid (wrong length, bad area code, etc.).
    """
    digits = _DIGIT_RE.sub('', raw)

    # Strip leading country code  (+1 / 1)
    if len(digits) == 11 and digits[0] == '1':
        digits = digits[1:]

    if len(digits) != 10:
        return None

    if not _VALID_RE.match(digits):
        return None

    return digits


def format_number(num: str) -> str:
    """'9032757138' → '(903) 275-7138'."""
    return f"({num[:3]}) {num[3:6]}-{num[6:]}"


def is_excel(filename: str) -> bool:
    return (filename or '').lower().endswith(EXCEL_EXTENSIONS)


# ── Cell helpers ──────────────────────────────────────────────────────────────

def cell_text(value: Any) -> str:
    """Render any cell value (str / int / float / date / None) as text."""
    if value is None:
        return ''
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, float):
        # Excel stores phone numbers typed into numeric cells as floats
        # (4405829719.0). Render them without the trailing '.0'.
        if value.is_integer():
            return str(int(value))
        return repr(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, datetime):
        if value.hour == 0 and value.minute == 0 and value.second == 0:
            return value.date().isoformat()
        return value.isoformat(sep=' ')
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _token_normalize(text: str) -> str | None:
    """normalize() with a cheap pre-filter for obviously non-phone tokens."""
    if not text:
        return None
    digit_count = sum(c.isdigit() for c in text)
    if digit_count < 7:
        return None
    return normalize(text)


# ── Raw row readers ───────────────────────────────────────────────────────────

def _decode(file_obj) -> str:
    """Read a binary file-like object as text (UTF-8, falling back to latin-1)."""
    data = file_obj.read()
    if isinstance(data, str):
        return data.lstrip('﻿')
    try:
        content = data.decode('utf-8')
    except UnicodeDecodeError:
        content = data.decode('latin-1')
    return content.lstrip('﻿')


def _sniff_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=',;\t|').delimiter
    except csv.Error:
        return ','


def iter_csv_rows(file_obj) -> Iterator[list[str]]:
    """Yield each non-empty row of a CSV/TXT file as a list of stripped strings."""
    content = _decode(file_obj)
    if not content:
        return
    delimiter = _sniff_delimiter(content[:65536])
    reader = csv.reader(io.StringIO(content, newline=''), delimiter=delimiter)
    for row in reader:
        cells = [c.strip() for c in row]
        if any(cells):
            yield cells


def iter_excel_rows(file_obj) -> Iterator[list[Any]]:
    """Yield each non-empty row of the first worksheet, keeping native cell values."""
    import openpyxl

    # openpyxl needs a seekable stream; Django storage files are not always one.
    data = file_obj.read()
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        for row in ws.iter_rows(values_only=True):
            cells = list(row)
            # Trim trailing empty cells Excel likes to pad rows with
            while cells and (cells[-1] is None or cells[-1] == ''):
                cells.pop()
            if any(c is not None and cell_text(c) != '' for c in cells):
                yield cells
    finally:
        wb.close()


def iter_rows(file_obj, filename: str = '') -> Iterator[list[Any]]:
    if is_excel(filename):
        return iter_excel_rows(file_obj)
    return iter_csv_rows(file_obj)


# ── Parsed representation ─────────────────────────────────────────────────────

@dataclass
class ParsedFile:
    """
    Everything the scrub pipeline needs from an upload.

    numbers:    unique valid 10-digit numbers, in order of first appearance
    rows:       number → the full original row (first occurrence wins)
    header:     the original header row, or None if the file had none
    phone_col:  index of the phone column (None if it could not be determined)
    width:      widest row seen (used to build a synthetic header)
    total_rows: data rows read (excluding header / blank rows)
    invalid:    rows without a valid US number
    duplicates: rows whose number was already seen
    """
    numbers:    list[str] = field(default_factory=list)
    rows:       dict[str, list[Any]] = field(default_factory=dict)
    header:     list[str] | None = None
    phone_col:  int | None = None
    width:      int = 0
    total_rows: int = 0
    invalid:    int = 0
    duplicates: int = 0

    @property
    def output_header(self) -> list[str]:
        """Header for result files: the original one, or a synthetic one."""
        if self.header:
            hdr = [cell_text(h) for h in self.header]
            # Pad in case data rows are wider than the header
            hdr += [f'column_{i + 1}' for i in range(len(hdr), self.width)]
            return hdr
        if self.width <= 1:
            return ['phone_number']
        hdr = [f'column_{i + 1}' for i in range(self.width)]
        if self.phone_col is not None and self.phone_col < len(hdr):
            hdr[self.phone_col] = 'phone_number'
        return hdr


def _looks_like_header(row: list[Any]) -> bool:
    """A first row is a header if it has words and no valid phone number."""
    texts = [cell_text(c) for c in row]
    if not any(_ALPHA_RE.search(t) for t in texts):
        return False
    return not any(_token_normalize(t) for t in texts)


def _phone_col_from_header(header: list[Any]) -> int | None:
    texts = [cell_text(h) for h in header]
    for i, h in enumerate(texts):
        if h.lower().replace(' ', '_') in ('phone_number', 'phone', 'cellphone', 'cell_phone', 'mobile'):
            return i
    for i, h in enumerate(texts):
        if _PHONE_HEADER_RE.search(h):
            return i
    return None


def _find_number(row: list[Any], preferred_col: int | None) -> tuple[str | None, int | None]:
    """
    Return (normalised_number, column_index) for a row.

    The preferred column (from the header or earlier rows) is tried first;
    otherwise every cell is scanned so files with the phone in an
    unexpected column still work.
    """
    if preferred_col is not None and preferred_col < len(row):
        norm = _token_normalize(cell_text(row[preferred_col]))
        if norm:
            return norm, preferred_col
    for i, cell in enumerate(row):
        if i == preferred_col:
            continue
        norm = _token_normalize(cell_text(cell))
        if norm:
            return norm, i
    return None, None


def parse_file(file_obj, filename: str = '') -> ParsedFile:
    """
    Fully parse an upload into a ParsedFile.

    Deduplication happens here so the task only pays the DNC check cost
    once per unique number; the first row carrying a number is the one
    that appears in the result files.
    """
    parsed = ParsedFile()
    seen = parsed.rows
    first = True
    phone_col: int | None = None

    for row in iter_rows(file_obj, filename):
        if first:
            first = False
            if _looks_like_header(row):
                parsed.header = row
                phone_col = _phone_col_from_header(row)
                parsed.phone_col = phone_col
                parsed.width = max(parsed.width, len(row))
                continue

        parsed.total_rows += 1
        parsed.width = max(parsed.width, len(row))

        norm, col = _find_number(row, phone_col)
        if norm is None:
            parsed.invalid += 1
            continue
        if phone_col is None:
            phone_col = col
            parsed.phone_col = col
        if norm in seen:
            parsed.duplicates += 1
            continue
        seen[norm] = row
        parsed.numbers.append(norm)

    return parsed


# ── Backwards-compatible helpers ──────────────────────────────────────────────

def extract_unique_numbers(file_obj, filename: str = '') -> tuple[list[str], int, int]:
    """
    Return (unique_valid_numbers, total_rows_read, invalid_count).

    Thin wrapper over parse_file() kept for callers that only need numbers.
    """
    parsed = parse_file(file_obj, filename)
    return parsed.numbers, parsed.total_rows, parsed.invalid
